import io
import os
import getpass
from typing import List, Optional

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import HTMLResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from .nri_client import AuthError, NRIClient

app = FastAPI(title="NRI Projects API", version="1.0.0")

client: Optional[NRIClient] = None

ENV_FILE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), ".env")


def load_env_file(path: str = ENV_FILE):
    """Читает NRI_USER / NRI_PASSWORD из .env (KEY=VALUE построчно)."""
    if not os.path.exists(path):
        return
    with open(path, encoding="utf-8-sig") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, value = line.partition("=")
            key, value = key.strip(), value.strip().strip('"').strip("'")
            if key in ("NRI_USER", "NRI_PASSWORD") and value:
                os.environ.setdefault(key, value)


def get_client() -> NRIClient:
    global client
    if client is None:
        try:
            client = NRIClient(
                os.environ.get("NRI_USER") or input("Логин NRI: "),
                os.environ.get("NRI_PASSWORD") or getpass.getpass("Пароль NRI: "),
            )
        except (EOFError, OSError):
            raise HTTPException(503, "Клиент не авторизован. Вызовите POST /api/login.")
    return client


@app.on_event("startup")
def startup():
    global client
    load_env_file()
    if os.environ.get("NRI_USER") and os.environ.get("NRI_PASSWORD"):
        try:
            c = NRIClient(os.environ["NRI_USER"], os.environ["NRI_PASSWORD"])
            c.login()
            client = c
            print(f"[NRI] Успешный вход: {os.environ['NRI_USER']} (из .env / переменных окружения)")
        except AuthError as e:
            print(f"[NRI] Вход не удался: {e}")
    else:
        print(f"[NRI] Задайте NRI_USER и NRI_PASSWORD в файле {ENV_FILE} "
              "или вызовите POST /api/login")


def _handle(e: Exception):
    if isinstance(e, AuthError):
        raise HTTPException(503, str(e))
    raise HTTPException(502, f"Ошибка NRI: {e}")


# ---------- API ----------

class LoginIn(BaseModel):
    username: str
    password: str


@app.post("/api/login")
def login(body: LoginIn):
    global client
    try:
        c = NRIClient(body.username, body.password)
        c.login()
        client = c
        return {"ok": True, "userId": c.user_id}
    except Exception as e:  # noqa: BLE001
        _handle(e)


@app.get("/api/projects")
def search_projects(query: Optional[str] = None, search_by: str = "all",
                    start: int = 0, limit: int = 100,
                    region_id: Optional[int] = None,
                    subsidiary_id: Optional[int] = None):
    """Поиск проектов (с пагинацией: start/limit, по умолчанию 100).
    search_by: all|number|project_id|bsname|posname|gfk|...
    region_id / subsidiary_id — фильтры как на сайте."""
    try:
        items, total = get_client().search_projects(
            query, search_by, start, limit, region_id, subsidiary_id)
        return {"total": total, "items": items}
    except Exception as e:  # noqa: BLE001
        _handle(e)


@app.get("/api/regions")
def list_regions():
    try:
        return {"items": get_client().list_regions()}
    except Exception as e:  # noqa: BLE001
        _handle(e)


@app.get("/api/subsidiaries")
def list_subsidiaries(region_id: Optional[int] = None):
    """Филиалы региона (без region_id — все)."""
    try:
        return {"items": get_client().list_subsidiaries(region_id)}
    except Exception as e:  # noqa: BLE001
        _handle(e)


@app.get("/api/projects/{pid}")
def get_project(pid: int):
    """Данные проекта (включая additionalInfo — «Дополнительная информация»)."""
    try:
        return get_client().get_project_info(pid)
    except Exception as e:  # noqa: BLE001
        _handle(e)


class CommentIn(BaseModel):
    text: str
    itype: Optional[int] = None


@app.post("/api/projects/{pid}/comment")
def set_comment(pid: int, body: CommentIn):
    """Записать текст в «Дополнительную информацию» проекта и сохранить."""
    try:
        return get_client().set_comment(pid, body.text, body.itype)
    except Exception as e:  # noqa: BLE001
        _handle(e)


@app.get("/api/projects/{pid}/actions")
def list_actions(pid: int, process_id: int):
    """Доступные действия проекта (панель «Доступные действия» на сайте)."""
    try:
        return {"items": get_client().list_actions(pid, process_id)}
    except Exception as e:  # noqa: BLE001
        _handle(e)


class RunActionIn(BaseModel):
    action_url: str
    action_id: int
    process_id: int
    comment: str = ""


@app.post("/api/projects/{pid}/run-action")
def run_action(pid: int, body: RunActionIn):
    """Выполнить доступное действие workflow без диалога."""
    try:
        return get_client().run_action(
            pid, body.process_id, body.action_url, body.action_id, body.comment)
    except Exception as e:  # noqa: BLE001
        _handle(e)


@app.get("/api/projects/{pid}/files")
def list_files(pid: int):
    """Список документов проекта (terrabyte)."""
    try:
        return {"items": get_client().list_files(pid)}
    except Exception as e:  # noqa: BLE001
        _handle(e)


@app.get("/api/projects/{pid}/doctypes")
def get_doc_types(pid: int):
    """Доступные типы документов для проекта."""
    try:
        return {"items": get_client().get_doc_types(pid)}
    except Exception as e:  # noqa: BLE001
        _handle(e)


@app.post("/api/projects/{pid}/attach")
async def attach(pid: int, file: UploadFile = File(...),
                 doc_type_id: str = Form(...), description: str = Form("")):
    """Прикрепить документ к проекту."""
    try:
        content = await file.read()
        return get_client().attach_file(
            pid, file.filename, content, doc_type_id, description)
    except Exception as e:  # noqa: BLE001
        _handle(e)


@app.post("/api/batch")
async def batch(
    project_ids: str = Form(..., description="JSON-список ID проектов"),
    comment: Optional[str] = Form(None),
    doc_type_id: Optional[str] = Form(None),
    description: str = Form(""),
    files: List[UploadFile] = File(None),
):
    """Пакетная обработка: комментарий + документы для нескольких проектов."""
    import json as _json
    try:
        ids = _json.loads(project_ids)
        if not isinstance(ids, list) or not all(isinstance(i, int) for i in ids):
            raise ValueError("project_ids должен быть JSON-массивом чисел")
    except ValueError as e:
        raise HTTPException(422, f"Некорректный project_ids: {e}")

    results = {}
    c = get_client()
    for pid in ids:
        res = {}
        try:
            if comment is not None and comment.strip():
                res["comment"] = c.set_comment(pid, comment)
            if files:
                if not doc_type_id:
                    raise RuntimeError("Для прикрепления файлов укажите doc_type_id")
                for f in files:
                    content = await f.read()
                    res.setdefault("attachments", []).append(
                        c.attach_file(pid, f.filename, content,
                                      doc_type_id, description))
        except Exception as e:  # noqa: BLE001
            res["error"] = str(e)
        results[str(pid)] = res
    return {"results": results}


# ---------- Excel: список ID + комментарии ----------

ID_HEADERS = {"id", "ид", "ид проекта", "projectid", "project_id",
              "id проекта", "номер id", "проект"}
ERP_HEADERS = {"ерп", "erp", "номер позиции", "позиция", "номер ерп"}
COMMENT_HEADERS = {"комментарий", "comment", "текст", "доп. информация",
                   "дополнительная информация", "дополнительно",
                   "комментарий (доп. информация)"}


@app.post("/api/excel/parse")
async def parse_excel(file: UploadFile = File(...)):
    """Разбирает Excel: колонка 1 — ID проекта, колонка 2 — ЕРП,
    колонка 3 — комментарий. Идентификация строки: по ID (приоритет),
    при пустом ID — по ЕРП."""
    from openpyxl import load_workbook
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(422, f"Не удалось прочитать Excel: {e}")

    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)
            if any(v is not None and str(v).strip() for v in r)]
    if not rows:
        raise HTTPException(422, "Файл пуст")

    # Ищем строку заголовков в первых 5 строках
    header_idx, id_col, erp_col, com_col = None, None, None, None
    for i, row in enumerate(rows[:5]):
        for j, cell in enumerate(row):
            val = str(cell or "").strip().lower().rstrip(":")
            if val in ID_HEADERS and id_col is None:
                id_col, header_idx = j, i
            elif val in ERP_HEADERS and erp_col is None:
                erp_col, header_idx = j, i
            elif val in COMMENT_HEADERS and com_col is None:
                com_col, header_idx = j, i
    notes = []
    if id_col is None and erp_col is None and com_col is None:
        # Заголовки не найдены: 1-я колонка = ID, 2-я = ЕРП, 3-я = комментарий
        id_col, erp_col, com_col = 0, 1, 2
        notes.append("Заголовки не распознаны — приняты: 1-я колонка = ID, "
                     "2-я = ЕРП, 3-я = комментарий")
        data_rows = rows
    else:
        if id_col is None:
            id_col = 0
        if erp_col is None:
            erp_col = 1
        if com_col is None:
            com_col = max(id_col, erp_col) + 1
        data_rows = rows[header_idx + 1:]

    def cell(row, col):
        return str(row[col]).strip() if col < len(row) and row[col] is not None else ""

    items, errors = [], []
    for n, row in enumerate(data_rows, start=1):
        raw_id = cell(row, id_col)
        raw_erp = cell(row, erp_col)
        comment = cell(row, com_col)
        if not raw_id and not raw_erp:
            errors.append(f"Строка {n}: не заполнен ни ID, ни ЕРП")
            continue
        item = {"comment": comment}
        if raw_id:
            try:
                item["id"] = int(float(raw_id))
            except ValueError:
                errors.append(f"Строка {n}: ID '{raw_id}' не является числом")
                continue
        else:
            item["erp"] = raw_erp
        items.append(item)

    return {"items": items, "total": len(items),
            "errors": errors, "notes": notes, "sheet": ws.title}


class CommentItem(BaseModel):
    id: int
    comment: str


class BatchCommentsIn(BaseModel):
    items: List[CommentItem]


@app.post("/api/batch-comments")
def batch_comments(body: BatchCommentsIn):
    """Внести комментарий в каждый проект из списка (ID + свой текст)."""
    c = get_client()
    results = {}
    for it in body.items:
        try:
            c.set_comment(it.id, it.comment)
            results[str(it.id)] = {"ok": True}
        except Exception as e:  # noqa: BLE001
            results[str(it.id)] = {"error": str(e)}
    ok_cnt = sum(1 for v in results.values() if v.get("ok"))
    return {"ok": ok_cnt, "failed": len(results) - ok_cnt, "results": results}


# ---------- мини-интерфейс ----------

app.mount("/static", StaticFiles(directory="static"), name="static")


@app.get("/", response_class=HTMLResponse)
def index():
    with open("static/index.html", encoding="utf-8") as f:
        return HTMLResponse(f.read())
